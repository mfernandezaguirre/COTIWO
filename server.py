import os,io,csv,json,sqlite3,secrets,hashlib,hmac,re,zipfile
from datetime import datetime,timedelta,timezone
from pathlib import Path
from fastapi import FastAPI,Request,HTTPException,UploadFile,File,Form,Depends
from fastapi.responses import HTMLResponse,FileResponse,StreamingResponse,RedirectResponse,JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel,Field
from openpyxl import Workbook,load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit

ROOT=Path(__file__).resolve().parent
DB=Path(os.environ.get('COTIWO_DB',str(ROOT/'cotiwo_demo.sqlite3')))
UPLOAD=ROOT/'uploads'; UPLOAD.mkdir(exist_ok=True)
app=FastAPI(title='Cotiwo Demo Integrada')
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')
app.mount('/uploads',StaticFiles(directory=UPLOAD),name='uploads')

def conn():
 c=sqlite3.connect(DB);c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON');return c

def digest(password,salt): return hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt),170000).hex()
def now():return datetime.now(timezone.utc).isoformat()
def guard_tax(value):
 value=re.sub(r'\D','',value or '')
 if len(value)!=11:return False
 weights=[5,4,3,2,7,6,5,4,3,2]; v=11-sum(int(a)*b for a,b in zip(value,weights))%11
 return (0 if v==11 else 9 if v==10 else v)==int(value[-1])

def init():
 with conn() as c:
  c.executescript('''
  CREATE TABLE IF NOT EXISTS companies(id TEXT PRIMARY KEY,name TEXT,kind TEXT,subscription TEXT DEFAULT 'pending',trial_end TEXT,created TEXT);
  CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,email TEXT UNIQUE,name TEXT,pass_hash TEXT,salt TEXT,role TEXT,company_id TEXT,tax_id TEXT,FOREIGN KEY(company_id) REFERENCES companies(id));
  CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY,user_id TEXT,expires TEXT);
  CREATE TABLE IF NOT EXISTS products(id INTEGER PRIMARY KEY AUTOINCREMENT,company_id TEXT,sku TEXT,name TEXT,category TEXT,description TEXT,price REAL,stock INTEGER,min_qty INTEGER,lead_days INTEGER,status TEXT,image TEXT,updated TEXT,UNIQUE(company_id,sku));
  CREATE TABLE IF NOT EXISTS quotes(id INTEGER PRIMARY KEY AUTOINCREMENT,company_id TEXT,buyer_id TEXT,contact TEXT,phone TEXT,notes TEXT,total REAL,review INTEGER,state TEXT,created TEXT,details TEXT,kind TEXT);
  CREATE TABLE IF NOT EXISTS notifications(id INTEGER PRIMARY KEY AUTOINCREMENT,company_id TEXT,quote_id INTEGER,text TEXT,is_read INTEGER DEFAULT 0,created TEXT);
  ''')
  if c.execute('SELECT COUNT(*) FROM companies').fetchone()[0]:return
  for id,name,kind in [('atl','Atlántica Trading','goods'),('global','Global Freight Logistics','forwarder'),('mixta','Conexión Integral','both')]:
   c.execute('INSERT INTO companies VALUES(?,?,?,?,?,?)',(id,name,kind,'active', (datetime.now(timezone.utc)+timedelta(days=14)).isoformat(),now()))
  def add(email,password,name,role,co,tax=''):
   salt=secrets.token_hex(16);c.execute('INSERT INTO users VALUES(?,?,?,?,?,?,?,?)',(secrets.token_hex(10),email,name,digest(password,salt),salt,role,co,tax))
  add('empresa@atlantica.demo','Empresa2026!','Administración Atlántica','company','atl')
  add('empresa@global.demo','Empresa2026!','Administración Global','company','global')
  add('empresa@mixta.demo','Empresa2026!','Administración Mixta','company','mixta')
  add('cliente@atlantica.demo','Cliente2026!','Comprador de prueba','buyer','atl')
  for co in ['atl','mixta']:
   for i,(name,cat,price) in enumerate([('Motor trifásico 3HP','Motores',295),('Motor trifásico 5HP','Motores',395),('Rodamiento 6204 ZZ','Mecánica',8.9),('Sensor inductivo M18','Sensores',45),('Reflector LED 200W','Iluminación',78),('Tablero IP65','Eléctricos',185),('Cable industrial 4mm','Eléctricos',2.8),('Variador 5kW','Automatización',570),('Contactor 32A','Eléctricos',26),('Motorreductor 0.75kW','Motores',430),('Fuente 24V 10A','Automatización',90),('PLC compacto','Automatización',890),('Cinta de embalaje','Insumos',5.5),('Caja de herramientas','Herramientas',95),('Guantes industriales','Seguridad',7.5),('Luminaria 100W','Iluminación',63),('Relé térmico','Eléctricos',39),('Sensor capacitivo','Sensores',70),('Tornillo M12','Mecánica',0.7),('Polea industrial','Mecánica',52),('Válvula neumática','Neumática',68),('Amoladora industrial','Herramientas',125),('Cilindro neumático','Neumática',115),('Rodamiento 6306','Mecánica',14.5)]):
    c.execute('INSERT INTO products(company_id,sku,name,category,description,price,stock,min_qty,lead_days,status,image,updated) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(co,f'{co.upper()}-{100+i}',name,cat,f'{name}, presentación industrial para cotizar.',price, 20+i*11,1 if i%4 else 5,2+i%7,'Activo','',now()))

init()

def current(req:Request, required_role=None):
 # Separate HttpOnly sessions allow company and buyer portals to coexist in
 # the same browser; role is checked against the authenticated DB record.
 portal = required_role or req.headers.get('x-cotiwo-portal') or req.query_params.get('portal') or 'company'
 if portal not in ('buyer', 'company'):
  raise HTTPException(400,'Portal inválido')
 token=req.cookies.get('cotiwo_'+portal+'_session')
 if not token:raise HTTPException(401,'Iniciá sesión en el portal correspondiente')
 with conn() as c:
  user=c.execute('SELECT u.* FROM sessions s JOIN users u ON s.user_id=u.id WHERE s.token_hash=? AND s.expires>?',(hashlib.sha256(token.encode()).hexdigest(),now())).fetchone()
  if not user:raise HTTPException(401,'Sesión expirada')
  if user['role']!=portal:raise HTTPException(403,'El usuario no pertenece a este portal')
  return dict(user)
def only(role):
 def inner(req:Request):
  return current(req,role)
 return inner
@app.get('/')
def root():return FileResponse(ROOT/'static'/'empresa.html')
@app.get('/clientes.html')
def buyers():return FileResponse(ROOT/'static'/'clientes.html')
class RegisterCompany(BaseModel):
 name:str;slug:str;kind:str;email:str;password:str;owner:str;tax_id:str
@app.post('/api/register-company')
def register_company(body:RegisterCompany):
 slug=body.slug.strip().lower()
 if not re.fullmatch(r'[a-z0-9][a-z0-9-]{2,39}',slug):raise HTTPException(400,'Identificador: entre 3 y 40 caracteres, letras, números o guiones')
 if body.kind not in ('goods','forwarder','both'):raise HTTPException(400,'Actividad inválida')
 if len(body.password)<8:raise HTTPException(400,'La contraseña requiere al menos 8 caracteres')
 if not guard_tax(body.tax_id):raise HTTPException(400,'CUIT/CUIL inválido')
 if not re.match(r'^[^\s@]+@[^\s@]+\.[^\s@]+$',body.email):raise HTTPException(400,'Correo inválido')
 if not body.name.strip() or not body.owner.strip():raise HTTPException(400,'Completá nombre de empresa y representante')
 salt=secrets.token_hex(16)
 with conn() as c:
  if c.execute('SELECT 1 FROM companies WHERE id=?',(slug,)).fetchone():raise HTTPException(409,'Ese enlace empresarial ya está ocupado')
  if c.execute('SELECT 1 FROM users WHERE email=?',(body.email.lower().strip(),)).fetchone():raise HTTPException(409,'El correo ya está registrado')
  c.execute('INSERT INTO companies(id,name,kind,subscription,trial_end,created) VALUES(?,?,?,?,?,?)',(slug,body.name.strip(),body.kind,'pending',None,now()))
  c.execute('INSERT INTO users VALUES(?,?,?,?,?,?,?,?)',(secrets.token_hex(10),body.email.lower().strip(),body.owner.strip(),digest(body.password,salt),salt,'company',slug,re.sub(r'\D','',body.tax_id)))
 return {'ok':True,'company_id':slug,'subscription':'pending'}

@app.get('/api/public/empresa/{company_id}')
def company_info(company_id:str):
 with conn() as c:
  co=c.execute('SELECT id,name,kind,subscription FROM companies WHERE id=?',(company_id,)).fetchone()
  if not co:raise HTTPException(404,'Proveedor desconocido')
  return dict(co)
class Login(BaseModel):
 email:str;password:str
@app.post('/api/login')
def login(body:Login):
 with conn() as c:
  u=c.execute('SELECT * FROM users WHERE email=?',(body.email.lower().strip(),)).fetchone()
  if not u or not hmac.compare_digest(u['pass_hash'],digest(body.password,u['salt'])):raise HTTPException(401,'Correo o contraseña incorrectos')
  token=secrets.token_urlsafe(48)
  c.execute('INSERT INTO sessions VALUES(?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),u['id'],(datetime.now(timezone.utc)+timedelta(hours=12)).isoformat()))
  response=JSONResponse({'role':u['role'],'company_id':u['company_id'],'name':u['name']})
  response.set_cookie('cotiwo_'+u['role']+'_session',token,httponly=True,samesite='lax',secure=os.environ.get('COOKIE_SECURE','false').lower()=='true',max_age=43200)
  return response
@app.post('/api/logout')
def logout(req:Request):
 portal=req.headers.get('x-cotiwo-portal') or req.query_params.get('portal') or 'company'
 if portal not in ('buyer','company'):raise HTTPException(400,'Portal inválido')
 tok=req.cookies.get('cotiwo_'+portal+'_session')
 if tok:
  with conn() as c:c.execute('DELETE FROM sessions WHERE token_hash=?',(hashlib.sha256(tok.encode()).hexdigest(),))
 resp=JSONResponse({'ok':True});resp.delete_cookie('cotiwo_'+portal+'_session');return resp
@app.get('/api/me')
def me(req:Request):
 u=current(req)
 with conn() as c:co=c.execute('SELECT id,name,kind,subscription,trial_end FROM companies WHERE id=?',(u['company_id'],)).fetchone()
 return {'email':u['email'],'name':u['name'],'role':u['role'],'tax_id':u['tax_id'],'company':dict(co)}
class Register(BaseModel):
 company_id:str;name:str;email:str;password:str;tax_id:str
@app.post('/api/register-buyer')
def register(body:Register):
 if not guard_tax(body.tax_id):raise HTTPException(400,'CUIT/CUIL inválido: verificá los 11 dígitos')
 if len(body.password)<8:raise HTTPException(400,'La contraseña debe tener al menos 8 caracteres')
 if not re.match(r'^[^\s@]+@[^\s@]+\.[^\s@]+$',body.email):raise HTTPException(400,'Correo inválido')
 with conn() as c:
  co=c.execute('SELECT id,subscription FROM companies WHERE id=?',(body.company_id,)).fetchone()
  if not co:raise HTTPException(404,'Empresa inexistente')
  if co['subscription']!='active':raise HTTPException(403,'El portal todavía no está habilitado')
  salt=secrets.token_hex(16)
  try:c.execute('INSERT INTO users VALUES(?,?,?,?,?,?,?,?)',(secrets.token_hex(10),body.email.lower().strip(),body.name,digest(body.password,salt),salt,'buyer',body.company_id,re.sub(r'\D','',body.tax_id)))
  except sqlite3.IntegrityError:raise HTTPException(409,'Ya existe una cuenta con ese correo')
 return {'ok':True}
@app.get('/api/products')
def list_products(req:Request):
 u=current(req)
 with conn() as c:
  rows=c.execute('SELECT * FROM products WHERE company_id=? ORDER BY updated DESC, id DESC',(u['company_id'],)).fetchall()
 return [dict(r) for r in rows] if u['role']=='company' else [{k:v for k,v in dict(r).items() if k!='price'} for r in rows if r['status']=='Activo']
class Product(BaseModel):
 sku:str;name:str;category:str='General';description:str='';price:float=Field(ge=0);stock:int=Field(ge=0);min_qty:int=Field(ge=1);lead_days:int=Field(ge=0);status:str='Activo'
@app.post('/api/products')
def save_product(body:Product,u=Depends(only('company'))):
 if not body.sku.strip() or not body.name.strip():raise HTTPException(400,'SKU y nombre obligatorios')
 with conn() as c:
  c.execute('''INSERT INTO products(company_id,sku,name,category,description,price,stock,min_qty,lead_days,status,image,updated) VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(company_id,sku) DO UPDATE SET name=excluded.name,category=excluded.category,description=excluded.description,price=excluded.price,stock=excluded.stock,min_qty=excluded.min_qty,lead_days=excluded.lead_days,status=excluded.status,updated=excluded.updated''',(u['company_id'],body.sku.strip(),body.name.strip(),body.category,body.description,body.price,body.stock,body.min_qty,body.lead_days,body.status,'',now()))
 return {'ok':True}
@app.post('/api/products/photo/{sku}')
async def photo(sku:str,image:UploadFile=File(...),u=Depends(only('company'))):
 raw=await image.read()
 if len(raw)>5_000_000:raise HTTPException(400,'Máximo 5MB')
 from PIL import Image
 try:
  im=Image.open(io.BytesIO(raw));im.thumbnail((1100,1100));im=im.convert('RGB')
 except Exception:raise HTTPException(400,'Imagen inválida')
 file=f'{u["company_id"]}_{secrets.token_hex(8)}.jpg'; im.save(UPLOAD/file,quality=82)
 with conn() as c:
  updated=c.execute('UPDATE products SET image=?,updated=? WHERE company_id=? AND sku=?',(f'/uploads/{file}',now(),u['company_id'],sku)).rowcount
  if not updated:raise HTTPException(404,'Producto no encontrado')
 return {'ok':True,'image':f'/uploads/{file}'}
@app.post('/api/products/import')
async def import_products(file:UploadFile=File(...),u=Depends(only('company'))):
 raw=await file.read()
 if len(raw)>3_000_000:raise HTTPException(400,'Archivo demasiado grande')
 if file.filename.lower().endswith('.xlsx'):
  wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
  rows=list(wb.active.values)
 elif file.filename.lower().endswith('.csv'):
  text=raw.decode('utf-8-sig');rows=list(csv.reader(io.StringIO(text),delimiter=';' if text.splitlines()[0].count(';')>text.splitlines()[0].count(',') else ','))
 else:raise HTTPException(400,'Usá CSV o XLSX')
 if not rows:return {'count':0,'errors':[]}
 headers=[str(h or '').strip().lower() for h in rows[0]];required={'sku','nombre','precio','stock'}
 if not required.issubset(headers):raise HTTPException(400,'Columnas requeridas: SKU, Nombre, Precio, Stock')
 errors=[];count=0
 with conn() as c:
  for i,r in enumerate(rows[1:],2):
   if not any(v is not None and str(v).strip() for v in r):continue
   x=dict(zip(headers,r))
   try:
    prod=Product(sku=str(x.get('sku') or '').strip(),name=str(x.get('nombre') or '').strip(),category=str(x.get('categoria') or 'General'),description=str(x.get('descripcion') or ''),price=float(x['precio']),stock=int(x['stock']),min_qty=int(x.get('minimo') or 1),lead_days=int(x.get('plazo_dias') or 3),status=str(x.get('estado') or 'Activo'))
    c.execute('''INSERT INTO products(company_id,sku,name,category,description,price,stock,min_qty,lead_days,status,image,updated) VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(company_id,sku) DO UPDATE SET name=excluded.name,category=excluded.category,description=excluded.description,price=excluded.price,stock=excluded.stock,min_qty=excluded.min_qty,lead_days=excluded.lead_days,status=excluded.status,updated=excluded.updated''',(u['company_id'],prod.sku,prod.name,prod.category,prod.description,prod.price,prod.stock,prod.min_qty,prod.lead_days,prod.status,'',now()));count+=1
   except Exception as e:errors.append(f'Fila {i}: {str(e)[:90]}')
 return {'count':count,'errors':errors[:30]}
@app.get('/api/products/export/{format}')
def export_products(format:str,u=Depends(only('company'))):
 with conn() as c: rows=[dict(r) for r in c.execute('SELECT sku,name,category,price,stock,min_qty,lead_days,status FROM products WHERE company_id=? ORDER BY sku',(u['company_id'],))]
 keys=['sku','nombre','categoria','precio','stock','minimo','plazo_dias','estado'];vals=lambda x:[x['sku'],x['name'],x['category'],x['price'],x['stock'],x['min_qty'],x['lead_days'],x['status']]
 b=io.BytesIO()
 if format=='xlsx':
  wb=Workbook();ws=wb.active;ws.title='Productos';ws.append(keys)
  for r in rows:ws.append(vals(r))
  wb.save(b);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 elif format=='csv':
  t=io.StringIO();w=csv.writer(t);w.writerow(keys)
  for r in rows:w.writerow(vals(r))
  b.write(('\ufeff'+t.getvalue()).encode());mime='text/csv'
 else:raise HTTPException(400,'Formato inválido')
 b.seek(0);return StreamingResponse(b,media_type=mime,headers={'Content-Disposition':f'attachment; filename="productos_cotiwo.{format}"'})
class Quote(BaseModel):
 items:list[dict]=[];phone:str='';notes:str='';kind:str='goods';logistics:dict={}
@app.post('/api/quotes')
def submit_quote(body:Quote,u=Depends(only('buyer'))):
 if not re.fullmatch(r'[+\d\s()\-]{6,25}',body.phone):raise HTTPException(400,'Teléfono inválido')
 with conn() as c:
  co=c.execute('SELECT * FROM companies WHERE id=?',(u['company_id'],)).fetchone()
  if co['subscription']!='active':raise HTTPException(403,'La licencia del proveedor no está activa')
  if co['kind']=='forwarder' or body.kind=='forwarder':
   if not body.logistics.get('origin') or not body.logistics.get('destination') or not body.logistics.get('weight'):raise HTTPException(400,'Completá origen, destino y peso')
   details=json.dumps({'logistics':body.logistics},ensure_ascii=False);total=0;review=1;kind='forwarder'
  else:
   if not body.items:raise HTTPException(400,'Seleccioná productos')
   total=0;review=0;details_list=[]
   # transactional stock decrement for submitted quotes
   for item in body.items:
    pid=int(item.get('id',0));qty=int(item.get('qty',0));p=c.execute('SELECT * FROM products WHERE id=? AND company_id=? AND status=?',(pid,u['company_id'],'Activo')).fetchone()
    if not p:raise HTTPException(404,'Producto no disponible')
    if qty < p['min_qty'] or qty > p['stock']:raise HTTPException(409,f'Cantidad inválida para {p["name"]} (mínimo {p["min_qty"]}; stock {p["stock"]})')
    total+=p['price']*qty;details_list.append({'sku':p['sku'],'name':p['name'],'qty':qty,'unit_price':p['price'],'subtotal':round(p['price']*qty,2)})
   for item in body.items:c.execute('UPDATE products SET stock=stock-?,updated=? WHERE id=? AND company_id=?',(int(item['qty']),now(),int(item['id']),u['company_id']))
   details=json.dumps({'items':details_list},ensure_ascii=False);kind='goods'
  contact=json.dumps({'name':u['name'],'email':u['email'],'tax_id':u['tax_id']})
  cursor=c.execute('INSERT INTO quotes(company_id,buyer_id,contact,phone,notes,total,review,state,created,details,kind) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(u['company_id'],u['id'],contact,body.phone,body.notes,total,review,'Nueva',now(),details,kind))
  qid=cursor.lastrowid
  c.execute('INSERT INTO notifications(company_id,quote_id,text,created) VALUES(?,?,?,?)',(u['company_id'],qid,f'Nueva solicitud #{qid} de {u["name"]} · {"requiere revisión" if review else "USD "+str(round(total,2))}',now()))
 return {'id':qid,'total':round(total,2),'review':bool(review),'state':'Nueva'}
@app.get('/api/quotes')
def quotes(req:Request):
 u=current(req)
 with conn() as c:
  where='company_id=?' if u['role']=='company' else 'company_id=? AND buyer_id=?';params=[u['company_id']] + ([] if u['role']=='company' else [u['id']])
  rows=c.execute('SELECT * FROM quotes WHERE '+where+' ORDER BY id DESC',params)
  return [{**dict(r),'details':json.loads(r['details']),'contact':json.loads(r['contact'])} for r in rows]
class Status(BaseModel):state:str
@app.patch('/api/quotes/{qid}')
def status(qid:int,body:Status,u=Depends(only('company'))):
 if body.state not in ['Nueva','En revisión','Enviada','Aceptada','Rechazada']:raise HTTPException(400,'Estado inválido')
 with conn() as c:
  if not c.execute('UPDATE quotes SET state=? WHERE id=? AND company_id=?',(body.state,qid,u['company_id'])).rowcount:raise HTTPException(404,'No encontrado')
 return {'ok':True}
@app.get('/api/notifications')
def notifications(u=Depends(only('company'))):
 with conn() as c:return [dict(r) for r in c.execute('SELECT * FROM notifications WHERE company_id=? ORDER BY id DESC LIMIT 60',(u['company_id'],))]
@app.post('/api/notifications/read')
def read_notifications(u=Depends(only('company'))):
 with conn() as c:c.execute('UPDATE notifications SET is_read=1 WHERE company_id=?',(u['company_id'],))
 return {'ok':True}
@app.get('/api/quotes/{qid}/pdf')
def quote_pdf(qid:int,req:Request):
 u=current(req)
 with conn() as c:
  q=c.execute('SELECT * FROM quotes WHERE id=? AND company_id=?'+(' AND buyer_id=?' if u['role']=='buyer' else ''),([qid,u['company_id'],u['id']] if u['role']=='buyer' else [qid,u['company_id']])).fetchone()
  if not q:raise HTTPException(404,'Sin acceso')
  company=c.execute('SELECT * FROM companies WHERE id=?',(u['company_id'],)).fetchone()
 out=io.BytesIO();p=canvas.Canvas(out,pagesize=A4);W,H=A4;y=H-65
 p.setFont('Helvetica-Bold',22);p.drawString(45,y,'COTIWO - PRESUPUESTO');y-=40;p.setFont('Helvetica',11)
 lines=[f'Proveedor: {company["name"]}',f'Solicitud: #{qid} | Estado: {q["state"]}',f'Fecha: {q["created"][:10]}',f'Cliente: {json.loads(q["contact"])["name"]}',f'Correo: {json.loads(q["contact"])["email"]}',f'CUIT/CUIL: {json.loads(q["contact"])["tax_id"]}','']
 d=json.loads(q['details'])
 if q['kind']=='goods':
  for it in d['items']:lines += [f'{it["sku"]} | {it["name"]} | {it["qty"]} u. | USD {it["subtotal"]:.2f}']
  lines += ['',f'TOTAL ESTIMADO: USD {q["total"]:.2f}']
 else:lines += [f'{k}: {v}' for k,v in d['logistics'].items()]+['','REQUIERE REVISION DE TARIFA']
 lines+=['','Documento orientativo. Sujeto a confirmacion comercial del proveedor.']
 for line in lines:
  for s in simpleSplit(str(line),'Helvetica',10,W-90):
   if y<55:p.showPage();y=H-65
   p.drawString(45,y,s);y-=19
 p.save();out.seek(0)
 return StreamingResponse(out,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="cotiwo_presupuesto_{qid}.pdf"'})
@app.get('/api/company')
def enterprise(u=Depends(only('company'))):
 with conn() as c:
  co=c.execute('SELECT * FROM companies WHERE id=?',(u['company_id'],)).fetchone();return dict(co)
@app.post('/api/demo/activate')
def activate(u=Depends(only('company'))):
 with conn() as c:c.execute('UPDATE companies SET subscription=? WHERE id=?',('active',u['company_id']))
 return {'ok':True,'note':'Suscripción de US$ 70/mes simulada, sin cobro'}
@app.post('/api/products/photos-batch')
async def photos_batch(file:UploadFile=File(...),u=Depends(only('company'))):
 raw=await file.read()
 if len(raw)>35_000_000:raise HTTPException(400,'Archivo ZIP demasiado grande (máximo 35 MB)')
 try:z=zipfile.ZipFile(io.BytesIO(raw))
 except Exception:raise HTTPException(400,'Subí un ZIP válido')
 from PIL import Image
 count=0;missing=[];errors=[]
 with conn() as c:
  for info in z.infolist():
   if info.is_dir():continue
   sku=Path(info.filename).stem.strip()
   if not info.filename.lower().endswith(('.jpg','.jpeg','.png','.webp')):continue
   if info.file_size>5_000_000:errors.append(sku+': tamaño superior a 5 MB');continue
   p=c.execute('SELECT id FROM products WHERE company_id=? AND sku=?',(u['company_id'],sku)).fetchone()
   if not p:missing.append(sku);continue
   try:
    im=Image.open(io.BytesIO(z.read(info)));im.thumbnail((1100,1100));im=im.convert('RGB')
    path=f'{u["company_id"]}_{secrets.token_hex(8)}.jpg';im.save(UPLOAD/path,quality=82)
    c.execute('UPDATE products SET image=?,updated=? WHERE id=? AND company_id=?',(f'/uploads/{path}',now(),p['id'],u['company_id']));count+=1
   except Exception:errors.append(sku+': imagen inválida')
 return {'count':count,'not_found':missing[:30],'errors':errors[:30]}
